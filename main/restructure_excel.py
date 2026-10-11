"""從各 fold 的 run_metrics.json 產生 long-format 多 sheet 統計 Excel。

新版資料流（2026-07 起）：
  - 每個 fold 的 patch 指標 → {fold_dir}/run_metrics.json           （main_runner 寫）
  - 每個 recon            → {fold_dir}/post_test_*/run_metrics.json（main_runner 寫）
  - recon per-scene 明細  → {fold_dir}/post_test_*/gt_aware_recon/per_scene_iou.csv
  這個腳本掃 RESULT_BASE 下所有 fold dir，讀上述檔案，重建 DST（_restructured.xlsx）。

舊 wide-format Excel（LEGACY_XLSX）已淘汰，只在「一次性 migration」時被讀：
  凡是還沒有 run_metrics.json 的舊 fold dir，會從 LEGACY_XLSX 對應欄位補一份 json。
  全部補完後，本腳本就完全不需要舊 Excel 了（未來的 run 由 main_runner 直接寫 json）。

Sheet 結構（2026-07-12 重設計，Claude×Codex 共識）：
  🚦 決策 / 🗺 Scene Recon / 🔬 Patch診斷 / ⚙ Run設定
  - 決策指標 = scene-recon pooled_oil_iou（gate>0.362），置頂單獨一張，不混 patch。
  - patch 指標標「診斷用」；歷史 run 的 mIoU/F1-macro 曾是 oil-only，統一由 CM 重算真 macro。
  - 每跑完一 fold 由 main_runner 呼叫一次（全掃 idempotent + 原子替換）。
"""
from pathlib import Path
import os
import re
import csv
import json
import statistics
import yaml
from datetime import datetime
from collections import defaultdict

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# ---------- 路徑 ----------
RESULT_BASE = Path('/mnt/backup/alanyh/oil_IR_Fullband/OIL_PROJECT_MutiBand_GT_expand/result-seg/CV_358clean_gt_expand')
LEGACY_XLSX = Path('/mnt/backup/alanyh/oil_IR_Fullband/OIL_PROJECT_MutiBand_GT_expand/result_excel/result_CV_358clean_gt_expand_log.xlsx')  # 僅供一次性 migration
DST = LEGACY_XLSX.with_name(LEGACY_XLSX.stem + '_restructured.xlsx')

# ---------- 決策常數 ----------
GATE = 0.362                     # 雙閘門 pooled_oil_iou 絕對門檻
EXPECTED_FOLDS = 3               # 本專案 CV 一律 3-fold；未滿即標「臨時」不出正式 PASS/FAIL
# 固定 baseline 參考（v2_stratified 三折 canonical，與 analysis/ 的 double_gate 腳本一致；
# 刻意不採「同資料夾上一個實驗」以免 baseline 漂移）。
BASELINE_POOLED_BY_FOLD = {'fold1': 0.3695, 'fold2': 0.3119, 'fold3': 0.3157}
BASELINE_POOLED_MEAN = sum(BASELINE_POOLED_BY_FOLD.values()) / len(BASELINE_POOLED_BY_FOLD)


def true_macro_from_cm(patch):
    """從 patch 混淆矩陣四格算含所有類別的真 macro。

    歷史 log 的 mIoU/F1-macro 曾誤存為 oil-only；新舊 run 一律由 CM 重算，
    避免欄位語意版本不同。
    回傳 dict(iou_oil, iou_bg, miou, f1_oil, f1_bg, f1_macro)；缺資料回 None。"""
    try:
        TP = float(patch['CM_GT(Oil)_Pred(Oil)'])
        FN = float(patch['CM_GT(Oil)_Pred(Background)'])
        FP = float(patch['CM_GT(Background)_Pred(Oil)'])
        TN = float(patch['CM_GT(Background)_Pred(Background)'])
    except (KeyError, TypeError, ValueError):
        return None

    def _safe(n, d):
        return n / d if d > 0 else None

    iou_oil = _safe(TP, TP + FP + FN)
    iou_bg = _safe(TN, TN + FP + FN)
    f1_oil = _safe(2 * TP, 2 * TP + FP + FN)
    f1_bg = _safe(2 * TN, 2 * TN + FN + FP)
    if None in (iou_oil, iou_bg, f1_oil, f1_bg):
        return None
    return dict(iou_oil=iou_oil, iou_bg=iou_bg, miou=(iou_oil + iou_bg) / 2,
                f1_oil=f1_oil, f1_bg=f1_bg, f1_macro=(f1_oil + f1_bg) / 2)

# ---------- run 標籤映射（依 timestamp prefix；僅供 legacy 分組/備註） ----------
RUN_TAGS = {
    '20260625-075607': ('legacy_baseline',   '5-fold CV 初版（無 C1-C7 fix, [13,1], no 255 ignore handling）'),
    '20260625-170038': ('legacy_baseline',   ''),
    '20260626-045600': ('legacy_baseline',   ''),
    '20260626-115614': ('legacy_baseline',   ''),
    '20260627-020537': ('legacy_baseline',   ''),
    '20260627-152751': ('C1-C7_A1_no_metric_fix', 'C1-C7 + A1 ImageNet pretrained，[13,1]，舊 eval 指標'),
    '20260627-220713': ('C1-C7_A1_no_metric_fix', ''),
    '20260628-042719': ('C1-C7_A1_no_metric_fix', ''),
    '20260628-162749': ('alpha_0.5_metric_fix',   '[0.5,1] + (A) metric fix + GT-aware recon'),
    '20260629-014016': ('balanced_with_bg_smoke', 'smoke: [1,1] + BG 補強 + (A) metric'),
    '20260629-052435': ('balanced_with_bg_smoke', 'smoke 第二版'),
    '20260629-065029': ('balanced_with_bg',       '[1,1] + BG 補強 + (A) metric'),
    '20260630-011954': ('with_regularization',    '加 Dropout(0.3) + WD=1e-3 + patience=25'),
    '20260630-054125': ('ignore_to_bg',           '加 255→bg，訓練有效 pixel 12%→100%（breakthrough）'),
    '20260630-130251': ('ignore_to_bg',           'fold2 of 3-fold CV'),
    '20260630-194912': ('ignore_to_bg',           'fold3 of 3-fold CV'),
    '20260701-133158': ('v2_stratified_3fold',    'v2 stratified 3-fold（每 fold test 不重疊、油汙大小均衡）fold1'),
    '20260701-204746': ('v2_stratified_3fold',    'v2 stratified 3-fold fold2'),
    '20260702-032032': ('v2_stratified_3fold',    'v2 stratified 3-fold fold3'),
}

# ---------- 手動指定的 CV group（override 自動偵測）----------
MANUAL_CV_GROUPS = [
    {
        'group_id': '20260630_ignore_to_bg_3fold',
        'tag':      'ignore_to_bg',
        'notes':    '3-fold CV，255→bg + BG 補強 + Dropout(0.3) + WD=1e-3 + patience=25',
        'folds': {
            'fold1': '20260630-054125',
            'fold2': '20260630-130251',
            'fold3': '20260630-194912',
        },
    },
]

FOLD_DIR_RE = re.compile(r'^(\d{8}-\d{6})_(.+)_(fold\d+)$')
LEGACY_COL_RE = re.compile(r'^(\d{8}-\d{6})_(.+?)_(fold\d+|CV_Summary)(_recon)?$')


# ---------- 工具 ----------
def f(v, ndigits=4):
    """安全格式化數字。"""
    if v is None or v == '':
        return ''
    try:
        return round(float(v), ndigits)
    except (ValueError, TypeError):
        return v


def load_json(p: Path):
    if not p.exists():
        return None
    try:
        with open(p, encoding='utf-8') as fh:
            return json.load(fh)
    except Exception:
        return None


def lookup_tag(ts):
    return RUN_TAGS.get(ts, ('unknown', ''))


def _corrected_oil_bg_ratio(fold_dir: Path, patch: dict):
    """val() 的 test 評估未依 scene 過濾背景 patch（deeplab_adapter.py:975 少傳
    bg_scene_filter），confusion matrix 混入了 bg_coord_txt 整份背景 patch（非僅
    此 fold 場景）。用 args.yaml 記錄的 bg_coord_txt 行數 × patch_size² 反推汙染
    像素量並扣除，還原這個 fold 真正的 GT-centric patch 油汙:背景像素比例。
    回傳 (corrected_oil_ratio, corrected_bg_ratio)，資料不足時回傳 (None, None)。"""
    args_path = fold_dir / 'args.yaml'
    if not args_path.exists():
        return None, None
    try:
        args_cfg = yaml.safe_load(args_path.read_text(encoding='utf-8')) or {}
    except Exception:
        return None, None
    ds_cfg = args_cfg.get('dataset') or {}
    bg_coord_txt = ds_cfg.get('bg_coord_txt')
    patch_size = int(ds_cfg.get('patch_size', 256))
    if not bg_coord_txt or not Path(bg_coord_txt).exists():
        return None, None
    try:
        with open(bg_coord_txt, encoding='utf-8') as fh:
            n_bg_total = sum(1 for _ in fh) - 1  # 扣 header
    except Exception:
        return None, None
    if n_bg_total <= 0:
        return None, None
    try:
        oil_px = float(patch.get('TestPixels(Oil)'))
        bg_px_all = float(patch.get('TestPixels(Background)'))
    except (TypeError, ValueError):
        return None, None
    contamination_px = n_bg_total * patch_size * patch_size
    clean_bg_px = bg_px_all - contamination_px
    if clean_bg_px <= 0:
        return None, None
    clean_total = oil_px + clean_bg_px
    if clean_total <= 0:
        return None, None
    return oil_px / clean_total, clean_bg_px / clean_total


# ---------- 一次性 migration：舊 wide Excel → fold dir 的 run_metrics.json ----------
def migrate_from_legacy():
    """凡是還沒有 run_metrics.json 的 fold dir / recon dir，從舊 Excel 對應欄位補一份。

    只補缺的（idempotent）；全部補完後舊 Excel 即淘汰。回傳補了幾份。
    """
    if not LEGACY_XLSX.exists():
        return 0
    wb = openpyxl.load_workbook(LEGACY_XLSX, data_only=True)
    ws = wb.active

    row_label = {}
    for r in range(1, ws.max_row + 1):
        v = ws.cell(row=r, column=1).value
        if v is not None:
            row_label[str(v).strip()] = r

    migrated = 0
    for c in range(2, ws.max_column + 1):
        col = str(ws.cell(row=1, column=c).value or '')
        m = LEGACY_COL_RE.match(col)
        if not m:
            continue
        ts, _base, fold, is_recon = m.group(1), m.group(2), m.group(3), bool(m.group(4))
        if fold == 'CV_Summary':
            continue  # CV summary 交給 restructure 自己重算

        # 抽出這一欄所有非空 (label -> value)
        d = {}
        for lab, r in row_label.items():
            if lab == 'Metric / Parameter':
                continue
            val = ws.cell(row=r, column=c).value
            if val is not None and val != '':
                d[lab] = val
        if not d:
            continue

        cands = sorted(RESULT_BASE.glob(f'{ts}_*_{fold}'))
        if not cands:
            continue
        fold_dir = cands[0]

        if is_recon:
            posts = sorted(fold_dir.glob('post_test_*'))
            for pt in posts:
                jp = pt / 'run_metrics.json'
                if not jp.exists():
                    dd = dict(d); dd['_kind'] = 'recon'
                    jp.write_text(json.dumps(dd, ensure_ascii=False, indent=2), encoding='utf-8')
                    migrated += 1
        else:
            jp = fold_dir / 'run_metrics.json'
            if not jp.exists():
                dd = dict(d); dd['_kind'] = 'patch'
                jp.write_text(json.dumps(dd, ensure_ascii=False, indent=2), encoding='utf-8')
                migrated += 1
    return migrated


# ---------- 掃 fold dir → merged ----------
def discover_runs():
    """回傳 {(ts, fold): {ts, fold, base, dir, patch, recon, pscsv}}。"""
    merged = {}
    for d in sorted(RESULT_BASE.glob('*_fold*')):
        if not d.is_dir():
            continue
        m = FOLD_DIR_RE.match(d.name)
        if not m:
            continue
        ts, base, fold = m.group(1), m.group(2), m.group(3)

        patch = load_json(d / 'run_metrics.json') or {}
        if patch:
            corr_oil, corr_bg = _corrected_oil_bg_ratio(d, patch)
            if corr_oil is not None:
                patch['TestPixelRatio(Oil)_corrected'] = f"{corr_oil:.4f}"
                patch['TestPixelRatio(Background)_corrected'] = f"{corr_bg:.4f}"

        recon = {}
        pscsv = None
        for pt in sorted(d.glob('post_test_*'), reverse=True):
            if not recon:
                rj = load_json(pt / 'run_metrics.json')
                if rj:
                    recon = rj
            if pscsv is None:
                cands = list(pt.rglob('per_scene_iou.csv'))
                if cands:
                    pscsv = cands[0]
        merged[(ts, fold)] = {
            'ts': ts, 'fold': fold, 'base': base, 'dir': d,
            'patch': patch, 'recon': recon, 'pscsv': pscsv,
        }
    return merged


def auto_detect_cv_groups(merged):
    """把「同 tag + 同 base + 有 ≥2 fold」的 run 歸為一組 CV group。"""
    groups = defaultdict(dict)
    for _key, m in merged.items():
        fold = m['fold']
        if not fold.startswith('fold'):
            continue
        tag = lookup_tag(m['ts'])[0]
        base = m.get('base', '')
        groups[(tag, base)][fold] = m['ts']

    auto = []
    for (tag, base), folds_dict in groups.items():
        fold_names = sorted(folds_dict.keys())
        if len(fold_names) < 2:
            continue
        first_ts = folds_dict[fold_names[0]]
        try:
            date_prefix = datetime.strptime(first_ts, '%Y%m%d-%H%M%S').strftime('%Y%m%d')
        except ValueError:
            date_prefix = first_ts.split('-')[0]
        auto.append({
            'group_id': f'{date_prefix}_{tag}_{len(fold_names)}fold_auto',
            'tag': tag,
            'notes': f'auto-detected {len(fold_names)}-fold CV (base={base})',
            'folds': folds_dict,
        })
    return auto


# ---------- 主函式 ----------
def main():
    # 路徑可由 CLI 覆寫（main_runner 會傳當前實驗的 results_base_dir / excel_log_path），
    # 不帶參數時維持舊行為（組0 CV_358clean_gt_expand）
    global RESULT_BASE, DST
    import argparse
    ap = argparse.ArgumentParser(description='從 run_metrics.json 重建 restructured Excel')
    ap.add_argument('--result-base', type=str, default=None,
                    help='要掃描的 result-seg 實驗目錄（預設=組0 CV_358clean_gt_expand）')
    ap.add_argument('--dst', type=str, default=None,
                    help='輸出 xlsx 路徑（預設=組0 的 _restructured.xlsx）')
    args = ap.parse_args()
    if args.result_base:
        RESULT_BASE = Path(args.result_base)
    if args.dst:
        DST = Path(args.dst)

    if not RESULT_BASE.exists():
        print(f'❌ result base 不存在：{RESULT_BASE}')
        return

    migrated = migrate_from_legacy()
    if migrated:
        print(f'[migration] 從舊 Excel 補了 {migrated} 份 run_metrics.json（一次性；補完即淘汰舊檔）')

    merged = discover_runs()
    sorted_keys = sorted(merged.keys(), key=lambda k: (k[0], k[1]))
    print(f'[掃描] {RESULT_BASE.name}：{len(sorted_keys)} 個 fold run')

    def pv(m, key, nd=4):
        return f(m['patch'].get(key), nd)

    def rv(m, key, nd=4):
        return f(m['recon'].get(key), nd)

    def recon_pooled(m):
        """該 fold 的 scene-recon pooled_oil_iou（決策指標）；無 recon 回 None。"""
        if not m['recon']:
            return None
        v = m['recon'].get('recon_gtaware_pooled_oil_iou')
        try:
            return float(v)
        except (TypeError, ValueError):
            return None

    # ---- 3-fold 平均用小工具 ----
    def group_members(grp):
        """回傳該 CV group 實際存在的 fold run（依 fold 名排序）。"""
        out = []
        for fold, ts in sorted(grp['folds'].items()):
            mm = merged.get((ts, fold))
            if mm:
                out.append(mm)
        return out

    def _fnum(v):
        try:
            return float(v)
        except (TypeError, ValueError):
            return None

    def _avg(vals):
        vals = [v for v in vals if v is not None]
        return round(sum(vals) / len(vals), 4) if vals else ''

    def _sum_int(vals):
        vals = [v for v in vals if v is not None]
        return int(sum(vals)) if vals else ''

    # ---- CV groups（決策摘要用）----
    manual_ts_sets = [frozenset(g['folds'].values()) for g in MANUAL_CV_GROUPS]
    auto_groups = auto_detect_cv_groups(merged)
    filtered_auto = [g for g in auto_groups if frozenset(g['folds'].values()) not in manual_ts_sets]
    all_groups = list(MANUAL_CV_GROUPS) + filtered_auto
    print(f'  [CV groups] manual={len(MANUAL_CV_GROUPS)}, auto={len(filtered_auto)}, total={len(all_groups)}')

    out_wb = openpyxl.Workbook()
    out_wb.remove(out_wb.active)

    def date_of(ts):
        try:
            return datetime.strptime(ts, '%Y%m%d-%H%M%S').strftime('%Y-%m-%d %H:%M')
        except ValueError:
            return ts

    # ======================================================================
    # Sheet 1: 🚦 決策（置頂）—— 只放 scene-recon pooled_oil_iou + gate + baseline delta
    # ======================================================================
    s_dec = out_wb.create_sheet('🚦 決策')
    s_dec.append(['決策指標 = Scene-recon Pooled Oil IoU（來自重建整張場景，不是 patch eval）。'
                  f'Gate: pooled > {GATE}'])
    s_dec.append([f'固定 baseline 參考: v2_stratified 三折 pooled = {BASELINE_POOLED_BY_FOLD}'
                  f' → mean {BASELINE_POOLED_MEAN:.4f}'])
    s_dec.append([])

    # -- CV 摘要表 --
    s_dec.append(['── CV 摘要（gate 判定用 = 各 fold pooled 的平均）──'])
    s_dec.append(['group', 'folds(done/exp)', 'CV mean pooled', 'Gate(>%.3f)' % GATE,
                  'gate margin', 'mean vs baseline Δ', 'avg best_val_miou'])
    for grp in all_groups:
        pooled_vals, bestvals = [], []
        for fname, ts in sorted(grp['folds'].items()):
            m = merged.get((ts, fname))
            if not m:
                continue
            p = recon_pooled(m)
            if p is not None:
                pooled_vals.append(p)
            bv = m['patch'].get('best_val_miou')
            if bv not in (None, ''):
                try:
                    bestvals.append(float(bv))
                except (TypeError, ValueError):
                    pass
        n_done = len(pooled_vals)
        if n_done == 0:
            continue
        mean_pooled = sum(pooled_vals) / n_done
        if n_done >= EXPECTED_FOLDS:
            gate = 'PASS' if mean_pooled > GATE else 'FAIL'
        else:
            gate = f'臨時({n_done}/{EXPECTED_FOLDS})'
        s_dec.append([grp['group_id'], f'{n_done}/{EXPECTED_FOLDS}',
                      round(mean_pooled, 4), gate, round(mean_pooled - GATE, 4),
                      round(mean_pooled - BASELINE_POOLED_MEAN, 4),
                      round(sum(bestvals) / len(bestvals), 4) if bestvals else ''])
    s_dec.append([])

    # -- 逐 fold 明細 --
    s_dec.append(['── 逐 fold 明細 ──'])
    s_dec.append(['run_id', 'fold', 'date', 'Recon pooled Oil IoU',
                  'Fold vs gate', 'vs baseline Δ', 'scenes', 'best_val_miou'])
    for key in sorted_keys:
        m = merged[key]
        fold = m['fold']
        pooled = recon_pooled(m)
        base_ref = BASELINE_POOLED_BY_FOLD.get(fold)
        if pooled is None:
            fold_gate, delta, pooled_cell = '', '', ''
        else:
            fold_gate = 'PASS' if pooled > GATE else 'FAIL'
            delta = round(pooled - base_ref, 4) if base_ref is not None else ''
            pooled_cell = round(pooled, 4)
        s_dec.append([f"{m['ts']}_{fold}", fold, date_of(m['ts']),
                      pooled_cell, fold_gate, delta,
                      rv(m, 'recon_gtaware_num_scenes', 0), pv(m, 'best_val_miou')])

    # ======================================================================
    # Sheet 2: 🗺 Scene Recon —— 只放 recon 指標，欄名全加 Recon 前綴
    # ======================================================================
    s_recon = out_wb.create_sheet('🗺 Scene Recon')
    s_recon.append([
        'run_id', 'Recon scene count',
        'Recon pooled Oil IoU', 'Recon pooled BG IoU',
        'Recon mean-per-scene Oil IoU', 'Recon mean mIoU',
    ])
    for key in sorted_keys:
        m = merged[key]
        if not m['recon']:
            continue
        s_recon.append([
            f"{m['ts']}_{m['fold']}", rv(m, 'recon_gtaware_num_scenes', 0),
            rv(m, 'recon_gtaware_pooled_oil_iou'), rv(m, 'recon_gtaware_pooled_bg_iou'),
            rv(m, 'recon_gtaware_mean_oil_iou'), rv(m, 'recon_gtaware_mean_miou'),
        ])
    # -- 每組 3-fold 平均 --
    s_recon.append([])
    s_recon.append(['── 3-fold 平均（每組 CV；scene count 為加總）──'])
    for grp in all_groups:
        mem = [mm for mm in group_members(grp) if mm['recon']]
        if len(mem) < 2:
            continue
        s_recon.append([
            f"▶ {grp['group_id']} (avg n={len(mem)})",
            _sum_int([_fnum(mm['recon'].get('recon_gtaware_num_scenes')) for mm in mem]),
            _avg([_fnum(mm['recon'].get('recon_gtaware_pooled_oil_iou')) for mm in mem]),
            _avg([_fnum(mm['recon'].get('recon_gtaware_pooled_bg_iou')) for mm in mem]),
            _avg([_fnum(mm['recon'].get('recon_gtaware_mean_oil_iou')) for mm in mem]),
            _avg([_fnum(mm['recon'].get('recon_gtaware_mean_miou')) for mm in mem]),
        ])

    # ======================================================================
    # Sheet 3: 🔬 Patch診斷 —— 診斷用，不作 gate 判定；mIoU/F1-macro 由 CM 重算含背景真 macro
    # ======================================================================
    s_patch = out_wb.create_sheet('🔬 Patch診斷')
    s_patch.append([
        'run_id',
        'Patch IoU(Oil)', 'Patch F1(Oil)', 'Patch Precision(Oil)', 'Patch Recall(Oil)',
        'Patch true mIoU', 'Patch true F1-macro', 'Patch Accuracy',
        'CM_TN', 'CM_FP', 'CM_FN', 'CM_TP',
        'TestPixelRatio(Oil)_corrected',
    ])
    for key in sorted_keys:
        m = merged[key]
        if not m['patch']:
            continue
        macro = true_macro_from_cm(m['patch'])
        s_patch.append([
            f"{m['ts']}_{m['fold']}",
            pv(m, 'IoU(Oil)'), pv(m, 'F1(Oil)'), pv(m, 'Precision(Oil)'), pv(m, 'Recall(Oil)'),
            round(macro['miou'], 4) if macro else '',
            round(macro['f1_macro'], 4) if macro else '',
            pv(m, 'Accuracy'),
            pv(m, 'CM_GT(Background)_Pred(Background)', 0), pv(m, 'CM_GT(Background)_Pred(Oil)', 0),
            pv(m, 'CM_GT(Oil)_Pred(Background)', 0), pv(m, 'CM_GT(Oil)_Pred(Oil)', 0),
            pv(m, 'TestPixelRatio(Oil)_corrected'),
        ])
    # -- 每組 3-fold 平均（rate 取平均、CM 四格取加總、true macro 由逐 fold 值平均）--
    s_patch.append([])
    s_patch.append(['── 3-fold 平均（每組 CV；CM 為加總，其餘取平均）──'])
    for grp in all_groups:
        mem = [mm for mm in group_members(grp) if mm['patch']]
        if len(mem) < 2:
            continue
        macros = [true_macro_from_cm(mm['patch']) for mm in mem]
        s_patch.append([
            f"▶ {grp['group_id']} (avg n={len(mem)})",
            _avg([_fnum(mm['patch'].get('IoU(Oil)')) for mm in mem]),
            _avg([_fnum(mm['patch'].get('F1(Oil)')) for mm in mem]),
            _avg([_fnum(mm['patch'].get('Precision(Oil)')) for mm in mem]),
            _avg([_fnum(mm['patch'].get('Recall(Oil)')) for mm in mem]),
            _avg([mac['miou'] for mac in macros if mac]),
            _avg([mac['f1_macro'] for mac in macros if mac]),
            _avg([_fnum(mm['patch'].get('Accuracy')) for mm in mem]),
            _sum_int([_fnum(mm['patch'].get('CM_GT(Background)_Pred(Background)')) for mm in mem]),
            _sum_int([_fnum(mm['patch'].get('CM_GT(Background)_Pred(Oil)')) for mm in mem]),
            _sum_int([_fnum(mm['patch'].get('CM_GT(Oil)_Pred(Background)')) for mm in mem]),
            _sum_int([_fnum(mm['patch'].get('CM_GT(Oil)_Pred(Oil)')) for mm in mem]),
            _avg([_fnum(mm['patch'].get('TestPixelRatio(Oil)_corrected')) for mm in mem]),
        ])

    # ======================================================================
    # Sheet 4: ⚙ Run設定 —— 合併訓練超參 + 資料統計
    # ======================================================================
    s_run = out_wb.create_sheet('⚙ Run設定')
    s_run.append([
        'run_id', 'date', 'architecture', 'seed', 'epochs', 'patience',
        'Batch size', 'Image size', 'lr0', 'use_amp', 'workers', 'training_time_minutes',
        'train_pos_patch', 'train_neg_patch', 'train_oil_ratio',
        'val_pos_patch', 'val_neg_patch', 'val_oil_ratio',
        'test_pos_patch', 'test_neg_patch', 'test_oil_ratio',
    ])
    for key in sorted_keys:
        m = merged[key]
        s_run.append([
            f"{m['ts']}_{m['fold']}", date_of(m['ts']),
            m['patch'].get('architecture') or '', pv(m, 'seed', 0), pv(m, 'Epochs', 0),
            pv(m, 'patience', 0), pv(m, 'Batch size', 0), pv(m, 'Image size', 0),
            pv(m, 'lr0', 6), m['patch'].get('use_amp') or '', pv(m, 'workers', 0),
            pv(m, 'training_time_minutes', 1),
            pv(m, 'train_pos_patch', 0), pv(m, 'train_neg_patch', 0), pv(m, 'train_oil_ratio'),
            pv(m, 'val_pos_patch', 0), pv(m, 'val_neg_patch', 0), pv(m, 'val_oil_ratio'),
            pv(m, 'test_pos_patch', 0), pv(m, 'test_neg_patch', 0), pv(m, 'test_oil_ratio'),
        ])

    # ---------- 樣式 ----------
    HEADER_FONT = Font(bold=True, color='FFFFFF')
    HEADER_FILL = PatternFill(start_color='3B5998', end_color='3B5998', fill_type='solid')
    SECTION_FONT = Font(bold=True, italic=True, color='333333')
    PASS_FILL = PatternFill(start_color='C6EFCE', end_color='C6EFCE', fill_type='solid')
    FAIL_FILL = PatternFill(start_color='FFC7CE', end_color='FFC7CE', fill_type='solid')
    PROV_FILL = PatternFill(start_color='FFEB9C', end_color='FFEB9C', fill_type='solid')
    AVG_FILL = PatternFill(start_color='DDE7F5', end_color='DDE7F5', fill_type='solid')
    border = Border(left=Side(style='thin', color='DDDDDD'),
                    right=Side(style='thin', color='DDDDDD'),
                    top=Side(style='thin', color='DDDDDD'),
                    bottom=Side(style='thin', color='DDDDDD'))
    HEADER_FIRSTS = {'run_id', 'group'}

    for ws in out_wb.worksheets:
        for row in ws.iter_rows():
            first = row[0].value
            is_header = isinstance(first, str) and first in HEADER_FIRSTS
            is_section = isinstance(first, str) and first.startswith('──')
            is_banner = isinstance(first, str) and (first.startswith('決策指標') or first.startswith('固定 baseline'))
            is_avg = isinstance(first, str) and first.startswith('▶')
            for cell in row:
                if is_header:
                    cell.font = HEADER_FONT
                    cell.fill = HEADER_FILL
                    cell.alignment = Alignment(horizontal='center', vertical='center')
                    cell.border = border
                elif is_avg:
                    cell.font = Font(bold=True)
                    cell.fill = AVG_FILL
                    cell.border = border
                elif is_section or is_banner:
                    cell.font = SECTION_FONT
                v = cell.value
                if v == 'PASS':
                    cell.fill = PASS_FILL
                elif v == 'FAIL':
                    cell.fill = FAIL_FILL
                elif isinstance(v, str) and v.startswith('臨時'):
                    cell.fill = PROV_FILL
        ws.freeze_panes = 'A2'
        for c in range(1, ws.max_column + 1):
            col_letter = get_column_letter(c)
            max_len = max((len(str(ws.cell(row=r, column=c).value or ''))
                           for r in range(1, ws.max_row + 1)), default=10)
            ws.column_dimensions[col_letter].width = min(max(max_len + 2, 12), 50)

    # ---------- 原子寫檔（研究者開著 Excel 時重產也不會壞檔）----------
    DST.parent.mkdir(parents=True, exist_ok=True)
    tmp = DST.with_name(DST.stem + '.tmp.xlsx')
    out_wb.save(tmp)
    os.replace(tmp, DST)
    print(f'\n✅ 寫入: {DST}')
    print(f'   {len(out_wb.worksheets)} sheets: {[s.title for s in out_wb.worksheets]}')
    for ws in out_wb.worksheets:
        print(f'   - {ws.title:<16} {ws.max_row} 列 × {ws.max_column} 欄')


if __name__ == '__main__':
    main()
